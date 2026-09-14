#!/usr/bin/env python3
"""The SIH 2026 IDEA-SUBMISSION deck: the official 6-slide format, exactly.

The idea round mandates six slides with fixed headers - TITLE PAGE, IDEA TITLE,
TECHNICAL APPROACH, FEASIBILITY AND VIABILITY, IMPACT AND BENEFITS, RESEARCH AND
REFERENCES - and penalises decks that ignore the format.  This generator fills
that skeleton with SAHYOG's measured content; the longer finale deck lives in
``scripts/make_win_deck.py`` and must never be submitted in its place.

    python3 scripts/make_idea_deck.py
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts._bootstrap import bootstrap  # noqa: E402

bootstrap()

from scripts.make_pitch_ppt import (AMBER, BG, CYAN, GREEN, INK, MUTED, PANEL,  # noqa: E402
                                    RED, add_rect, add_text, bullets, footer)
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "docs" / "SAHYOG_SIH2026_IDEA_Submission.pptx"

TEAM_NAME = "<TEAM NAME (as registered on portal)>"
TEAM_ID = "<TEAM ID>"
PS_TITLE = ("<paste exact title from the SIH portal> - autonomous UAV-based "
            "search, detection and assistance for disaster survivors")


def template_header(s, n: int, title: str):
    """The official template look: plain header, '@SIH Idea submission - Template'."""
    add_rect(s, 0, 0, 13.333, 7.5, BG)
    add_rect(s, 0, 0, 13.333, 1.05, PANEL)
    add_text(s, 0.55, 0.18, 9.5, 0.7, title, size=26, color=INK, bold=True)
    add_text(s, 0.55, 0.68, 9.5, 0.3, "@SIH Idea submission - Template", size=10,
             color=MUTED)
    add_text(s, 10.2, 0.3, 2.6, 0.5, TEAM_NAME, size=11, color=CYAN, bold=True)
    add_text(s, 12.6, 0.18, 0.4, 0.5, str(n), size=14, color=MUTED)


def s1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, 13.333, 7.5, BG)
    add_rect(s, 0.85, 1.0, 11.63, 0.06, GREEN)
    add_text(s, 0.85, 1.3, 11.6, 0.8, "TITLE PAGE", size=34, color=INK, bold=True)
    add_text(s, 0.85, 2.05, 11.6, 0.5, "SMART INDIA HACKATHON 2026", size=16,
             color=GREEN, bold=True)
    rows = [("Problem Statement ID", "SIH26177"),
            ("Problem Statement Title", PS_TITLE),
            ("Theme", "Robotics & Drones"),
            ("PS Category", "Hardware"),
            ("Organization", "Qualcomm Inc."),
            ("Team ID", TEAM_ID),
            ("Team Name (registered on portal)", TEAM_NAME),
            ("Institute (AISHE)", "<INSTITUTE NAME>  ·  <AISHE CODE>")]
    y = 2.85
    for k, v in rows:
        add_text(s, 0.85, y, 4.3, 0.42, k, size=13, color=MUTED, bold=True)
        add_text(s, 5.25, y, 7.2, 0.62, v, size=13.5, color=INK)
        y += 0.52
    return s, ("Idea submission, slide 1 of 6. Nothing to say here beyond the "
               "portal-exact problem statement title - judges check it matches.")


def s2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    template_header(s, 2, "IDEA TITLE")
    add_text(s, 0.55, 1.25, 12.2, 0.6,
             [("SAHYOG - the phone-powered thermal search drone",
               {"size": 22, "bold": True, "color": GREEN})])
    add_text(s, 0.55, 1.85, 12.2, 0.35, "IDEA / SOLUTION:", size=13, color=CYAN, bold=True)
    bullets(s, 0.55, 2.3, 12.2, 4.6, [
        ("The problem, in our words: ", "in floods and quakes, survivors are lost "
         "in the first hour because search is a human watching thermal video; GPS, "
         "networks and roads fail together; and at survey altitude a human body is "
         "4-10 pixels - smaller than one grid cell of an off-the-shelf detector "
         "(measured mAP50 0.0009, i.e. blind). Affected: flood/quake survivors and "
         "the SDRF/NDRF/district teams who reach them late."),
        ("The solution, in our words: ", "an autonomous drone that fuses a "
         "radiometric LWIR camera with an Android phone's RGB camera, detects "
         "survivors with a YOLO11n carrying an added stride-4 (P2) head, passes "
         "every detection through physics gates, and returns 200-byte verified "
         "survivor packets over LoRa - entirely on-device, no internet at any step."),
        ("The unfair advantage: ", "the phone supplies six subsystems (stabilised "
         "RGB, INT8 NPU, VIO, 4G relay, display, backup battery) = Rs 0 of a "
         "Rs 35,000+ payload; its Hexagon DSP is the same family as this PS's "
         "target QCS6490 board - one model, two destinations."),
        ("Who it serves: ", "district disaster control rooms; first-hour search of "
         "9 ha per 10-minute pass, night, smoke or canopy."),
    ], size=14, gap=8)
    return s, ("Problem first, then solution, then the advantage, then the "
               "beneficiary - in that order, one breath each. The 4-10 pixel "
               "number is the hook: it is why every other approach fails and why "
               "ours is shaped the way it is.")


def s3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    template_header(s, 3, "TECHNICAL APPROACH")
    bullets(s, 0.55, 1.3, 12.2, 5.6, [
        ("Detection: ", "YOLO11n + bundled P2 (stride-4) head - the single largest "
         "accuracy change for sub-8 px targets (0.0009 -> 0.234 mAP50, 260x); "
         "thermal trained with radiometric augmentation; ONNX opset 12 -> INT8 "
         "w8a8 QNN for Hexagon behind a >=97% recall-retention gate."),
        ("Physics gates (the trust layer): ", "radiometric test (body 30-37 C vs "
         "flood water 15 C; sun-heated roofs rejected), projected-extent geometry "
         "veto, two-pass confirmation (find at 45 m, confirm at 22 m), "
         "Platt-scaled confidences so the belief map composes probabilities."),
        ("Search strategy: ", "lawnmower lanes with spacing = sensor swath x "
         "detection-probability target (search theory); every negative observation "
         "lowers the belief map and drives re-visit."),
        ("Navigation: ", "EKF source-manager; phone VIO + lidar when GNSS is "
         "denied; position sigma monitored, sortie ends past 25 m."),
        ("Communications: ", "LoRa 900 MHz store-and-forward, 200-byte acked "
         "packets; ELRS is control-only so the alert path never depends on it."),
        ("Safety: ", "12-rule supervisor (25% reserve, geofence, wind, telemetry "
         "silence) outranks the mission."),
        ("Stack: ", "Python - ArduPilot/MAVLink - ONNX Runtime + QNN - ultralytics "
         "fine-tune on a free Kaggle T4 - datasets HIT-UAV / AIResQ / SARD."),
    ], size=13, gap=7)
    return s, ("Seven method lines, each ending in a number or a named algorithm. "
               "If a judge interrupts, any line can be defended from the "
               "repository. Do not read this slide; point at it.")


def s4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    template_header(s, 4, "FEASIBILITY AND VIABILITY")
    bullets(s, 0.55, 1.3, 12.2, 3.4, [
        ("Technical: ", "180+ automated tests pass; end-to-end simulated sorties "
         "today; phone bridge benchmarked on real hardware (15.0 fps RGB, 29.6 Hz "
         "VIO, 69 ms worst gap); ONNX + INT8 artefacts on disk; onboard budget "
         "165 ms vs 500 ms available; every layer has a fallback."),
        ("Financial: ", "phone replaces Rs 35,000+ of bought subsystems; airframe "
         "~Rs 40k; radiometric LWIR ~Rs 15-25k -> under Rs 1 L per aircraft vs "
         "Rs 3 L+ conventional; zero cloud cost, all inference on-device."),
        ("Market: ", "every district control room already owns the compute "
         "(phones); no new supply chain for the AI."),
        ("Operational: ", "one-operator launch; crew training is one Android app; "
         "maintenance is Android updates."),
    ], size=13, gap=7)
    add_rect(s, 0.55, 4.95, 12.2, 0.04, PANEL)
    add_text(s, 0.55, 5.1, 12.2, 0.35, "POTENTIAL CHALLENGES AND RISKS  ->  STRATEGIES FOR OVERCOMING",
             size=12, color=AMBER, bold=True)
    bullets(s, 0.55, 5.5, 12.2, 1.7, [
        ("LWIR lead time -> ", "Boson 320 second source.   "),
        ("Small-target recall gap -> ", "P2 head + 640x512 real-data fine-tune, "
         "already running on the free-GPU path.   "),
        ("INT8 accuracy loss -> ", "w8a16 fallback behind the same >=97% gate.   "),
        ("Monsoon window -> ", "dry-season trials in Bihar.   "),
        ("No local GPU -> ", "documented free-tier GPU runbook."),
    ], size=12, marker_color=AMBER, gap=3)
    return s, ("Four feasibility axes in the template's own order, then the risk "
               "table. Naming risks with owned mitigations is the single cheapest "
               "credibility win in this round.")


def s5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    template_header(s, 5, "IMPACT AND BENEFITS")
    add_text(s, 0.55, 1.3, 12.2, 0.35, "POTENTIAL IMPACT ON THE TARGET AUDIENCE",
             size=12, color=GREEN, bold=True)
    bullets(s, 0.55, 1.7, 12.2, 2.0, [
        ("Positive: ", "first-hour coverage at machine speed (9 ha per 10-minute "
         "pass); measured precision 1.000 - zero false alarms, so ground teams "
         "stop climbing wrong rooftops; LWIR works at night, through smoke, under "
         "canopy; verified reports survive total network loss."),
        ("Negative, named and mitigated: ", "technology-adoption risk -> "
         "phone-first design, one app, heuristic fallback so a missing model never "
         "grounds the aircraft; LWIR cost -> offset by the phone's Rs 35k "
         "substitution."),
    ], size=13.5, gap=7)
    add_text(s, 0.55, 3.95, 12.2, 0.35, "BENEFITS OF THE SOLUTION", size=12, color=CYAN, bold=True)
    bullets(s, 0.55, 4.35, 12.2, 2.6, [
        ("Social: ", "golden-hour rescues reached in the golden hour; rescuers "
         "routed around hazards (A* over the flood field); equity - "
         "district-affordable, not metro-affordable."),
        ("Economic: ", "Rs 0 compute payload; one airframe re-tasked across flood, "
         "wildfire, landslide and coastal SAR by swapping modality weights, not "
         "mission code."),
        ("Environmental: ", "electric airframe replaces helicopter search hours; "
         "minimal noise over distressed communities."),
    ], size=13.5, gap=7, marker_color=CYAN)
    return s, ("Answering the template's 'negative impact' prompt is the "
               "differentiator: ninety percent of teams pretend there are none. "
               "We name two and close both.")


def s6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    template_header(s, 6, "RESEARCH AND REFERENCES")
    bullets(s, 0.55, 1.3, 12.2, 3.6, [
        ("Datasets studied: ", "HIT-UAV (2,898 aerial LWIR frames) - AIResQ "
         "(airborne thermal SAR; source of the honest mAP50 ~ 0.55 ceiling) - SARD "
         "(1,981 SAR postures) - VisDrone - HERIDAL - RescueNet."),
        ("Literature: ", "MDPI J. Imaging 11(12):436 - 75k-image aerial-thermal "
         "benchmark: RT-DETR-L 5.0% vs YOLO 10.7-11.6% AP on small objects (why we "
         "fly a CNN, not a transformer) - Ultralytics YOLO11 benchmarks - Koopman "
         "search theory - Platt (1999) probability calibration."),
        ("Existing solutions and our difference: ", "commercial thermal drones = "
         "human-watched video, no autonomy, no rejection of false alarms; research "
         "UAV SAR = RGB-only or cloud-dependent; ours = on-device radiometric "
         "detection + physics gates + phone-borne compute, every claim traceable "
         "to a checksummed artefact in our repository."),
    ], size=13, gap=7)
    add_rect(s, 0.55, 5.0, 12.2, 0.04, PANEL)
    add_text(s, 0.55, 5.15, 12.2, 0.35, "TEAM", size=12, color=CYAN, bold=True)
    add_text(s, 0.55, 5.5, 12.2, 0.8,
             [("<Member 1> AI/perception  ·  <Member 2> flight & safety  ·  "
               "<Member 3> comms & GCS", {"size": 12.5, "color": INK})],
             line_spacing=1.2)
    add_text(s, 0.55, 5.95, 12.2, 0.8,
             [("<Member 4> mobile bridge  ·  <Member 5> airframe & power  ·  "
               "<Member 6> data & field ops", {"size": 12.5, "color": INK})],
             line_spacing=1.2)
    add_text(s, 0.55, 6.55, 12.2, 0.4,
             [("Contact: <email>   ·   Thank you - live demo available.",
               {"size": 13, "color": GREEN, "bold": True})])
    return s, ("Close on the difference line, then offer the demo. 'Every claim on "
               "these six slides traces to a checksummed artefact in our "
               "repository - we would rather show you the measurement than ask "
               "you to trust the adjective.'")


SLIDES = [s1, s2, s3, s4, s5, s6]


def build(out: Path = OUT) -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, fn in enumerate(SLIDES, 1):
        slide, note = fn(prs)
        slide.notes_slide.notes_text_frame.text = note
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return str(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", help="destination .pptx")
    args = ap.parse_args(argv)
    print(f"wrote {build(Path(args.out) if args.out else OUT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
